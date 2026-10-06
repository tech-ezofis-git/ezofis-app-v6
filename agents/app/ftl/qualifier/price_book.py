"""
FTL "Contractor Price Book" (.docx) ingestion — the structured sibling of the PDF-page indexer in
pricelist_store.reindex_pricelist().

Why this exists: the original UI upload expects the old 2026 pricelist PDF (one text blob per PDF
page, embedded as-is). FTL's REV. 09.2026 price book is a different animal: a designed Word
document whose catalogue sections are clean Word TABLES with an explicit "FTL ITEM CODE" column.
That is strictly better source material (no column-guessing from flattened PDF text), so this module
reads those tables directly and produces three things from ONE parse, so they can never drift apart:

  1. An *overlay* (data/rulebook/price_book_overlay.json) — canonical code -> price, plus the old
     codes each row replaced (aliases). rules_engine.py applies it on top of the Stage 2 Product
     Master so every deterministic price/code check in agent.py uses the price book, regardless of
     what the embedded index or the model says.
  2. Index *pages* (text) in the same "TYPE OEM CODE HAND DESCRIPTION WIDTH $ PRICE" row shape the
     existing regex parsers in pricelist_store.py were written against, so search_pricelist() and
     door_operator_prices()/governor_prices()/... keep working unchanged.
  3. A merge with the catalogue sections the price book does NOT cover (universal car door panels,
     VisionPlus detectors, OL100 governor, OL35 swingarm, spare parts): those are carried over
     verbatim from data/pricelist_legacy_blocks.json rather than silently dropped, so an estimate
     never loses a priceable line just because the new book is narrower than the old PDF.

Nothing here talks to OpenAI. Embedding the pages is pricelist_store's job (it needs an API key).
"""

from __future__ import annotations

import json
import os
import re
from typing import Any, Dict, List, Optional, Tuple

_QUALIFIER_DIR = os.path.dirname(__file__)
DATA_DIR = os.path.join(_QUALIFIER_DIR, "data")
LEGACY_BLOCKS_PATH = os.path.join(DATA_DIR, "pricelist_legacy_blocks.json")
# The overlay is applied by the shared rulebook loader, which lives with the Quote Estimator.
OVERLAY_PATH = os.path.join(
    os.path.dirname(_QUALIFIER_DIR), "quote_estimator", "data", "rulebook", "price_book_overlay.json"
)

# A page is embedded as chunk_text() chunks of 280 words. Keeping every generated page under this
# budget means each page is exactly one chunk, so a catalogue row can never be split across two
# chunks (the parsers in pricelist_store.py work chunk-by-chunk and would lose a split row's price).
PAGE_WORD_BUDGET = 260

_NOTE_PREFIXES = (
    "REQUIRED PAIRING",
    "PO REQUIREMENT",
    "SELECTION CHECKPOINT",
    "MAJOR PROJECTS",
    "SEND A COMPLETE REQUEST",
)
_PRICE_RE = re.compile(r"^\$?\s*([\d,]+\.\d{2})$")


class PriceBookError(ValueError):
    """The .docx is not (or no longer is) laid out like FTL's Contractor Price Book. Raised instead
    of guessing, because a half-parsed price book would silently price quotes wrong."""


# ---------------------------------------------------------------------------------------------
# Reading the .docx
# ---------------------------------------------------------------------------------------------
def _table_rows(table) -> List[List[str]]:
    """Cell text via raw XML <w:t> runs rather than python-docx's cell.text: cell.text silently
    drops text inside hyperlinks/fields, and this book is full of them (the 'click any product
    family' quick index)."""
    from docx.oxml.ns import qn

    rows: List[List[str]] = []
    for tr in table._tbl.findall(qn("w:tr")):
        cells = []
        for tc in tr.findall(qn("w:tc")):
            # Runs inside ONE paragraph are concatenated with no separator: Word freely splits a
            # single item code across several runs (confirmed: "SGV2(1S)_PANEL_ADAPTOR_MAC36_LH"
            # arrives as "SGV2(1", "S)_", "PANEL_ADAPTOR", ...), so joining runs with a space would
            # corrupt every such code. Only separate distinct paragraphs / explicit breaks.
            paras = []
            for p in tc.iter(qn("w:p")):
                buf = []
                for el in p.iter():
                    if el.tag == qn("w:t") and el.text:
                        buf.append(el.text)
                    elif el.tag in (qn("w:br"), qn("w:tab")):
                        buf.append(" ")
                paras.append("".join(buf))
            cells.append(re.sub(r"\s+", " ", " ".join(paras)).strip())
        rows.append(cells)
    return rows


def _open_docx(path: str):
    from docx import Document

    return Document(path)


def looks_like_price_book(path: str) -> bool:
    """True for a .docx with at least one 'FTL ITEM CODE' / 'CATALOGUE PRICE' table — the signature
    of the Contractor Price Book. Cheap and side-effect free; used by the UI to route uploads."""
    if not str(path).lower().endswith(".docx"):
        return False
    try:
        doc = _open_docx(path)
        for t in doc.tables:
            rows = _table_rows(t)
            if rows and "FTL ITEM CODE" in rows[0] and "CATALOGUE PRICE" in rows[0]:
                return True
    except Exception:
        return False
    return False


def _parse_price(text: str, where: str) -> float:
    m = _PRICE_RE.match((text or "").strip())
    if not m:
        raise PriceBookError(f"Could not read a price from {text!r} ({where}).")
    return float(m.group(1).replace(",", ""))


def classify_family(code: str) -> str:
    c = code.upper()
    if c.startswith("SGV2_DOOR_OP_"):
        return "door_operator"
    if c == "SGV2_DOOR_TOOLS":
        return "door_tools"
    if "_PANEL_ADAPTOR_" in c:
        return "panel_adaptor"
    if c.startswith("SGV2_CLUTCH_"):
        return "clutch"
    if c.startswith("WRG_"):
        return "roller_guide"
    if c.startswith("WG_"):
        return "governor"
    if c.startswith("WS_CSGB"):
        return "car_safety"
    raise PriceBookError(
        f"Unrecognized item code {code!r}: this price book contains a product family this importer "
        "does not know how to classify. Add it to price_book.classify_family() deliberately rather "
        "than letting it through unclassified."
    )


def derive_aliases(code: str, family: str) -> List[str]:
    """Codes this row REPLACED in the previous (PDF) pricelist / Stage 2 Product Master, derived by
    the two renames the REV. 09.2026 book made:
      - operators and clutches dropped the 'H' from the hand suffix (…_LH -> …_L, …_RH -> …_R);
      - panel adaptors renamed '_DP_' to '_PANEL_ADAPTOR_'.
    Old codes stay *accepted* (never rejected as fabricated) and are rewritten to the new code."""
    aliases: List[str] = []
    if family in ("door_operator", "clutch") and re.search(r"_(L|R)$", code):
        aliases.append(code + "H")
    if family == "panel_adaptor":
        aliases.append(code.replace("_PANEL_ADAPTOR_", "_DP_"))
    return aliases


def parse_price_book(path: str) -> Dict[str, Any]:
    """Parse the Contractor Price Book into {revision, rows, incentives, terms, notes}.

    Strict on purpose: duplicate codes, unreadable prices, or a missing catalogue table abort the
    import (PriceBookError) instead of producing a partial price list."""
    doc = _open_docx(path)

    revision = ""
    for t in doc.tables[:2]:
        for row in _table_rows(t):
            for cell in row:
                m = re.search(r"REV\.?\s*([0-9]{2}\.[0-9]{4})", cell, re.IGNORECASE)
                if m:
                    revision = f"REV. {m.group(1)}"
                    break
            if revision:
                break
        if revision:
            break

    rows_out: List[Dict[str, Any]] = []
    incentives: List[Dict[str, str]] = []
    terms: List[Dict[str, str]] = []
    notes: List[str] = []
    seen_codes: Dict[str, int] = {}

    for t in doc.tables:
        rows = _table_rows(t)
        if not rows:
            continue
        header = rows[0]

        if "FTL ITEM CODE" in header and "CATALOGUE PRICE" in header:
            col = {name: header.index(name) for name in header if name}
            desc_col = next(
                (n for n in ("CONFIGURATION", "COMPONENT", "DESCRIPTION") if n in col), None
            )
            for r in rows[1:]:
                if len(r) < len(header) or not any(r):
                    continue
                code = r[col["FTL ITEM CODE"]].strip()
                if not code:
                    continue
                price = _parse_price(r[col["CATALOGUE PRICE"]], f"code {code}")
                if code in seen_codes:
                    raise PriceBookError(f"Duplicate item code {code!r} in the price book.")
                seen_codes[code] = 1
                family = classify_family(code)
                width_txt = r[col["WIDTH"]] if "WIDTH" in col else ""
                width_m = re.search(r"\d+", width_txt or "")
                rows_out.append(
                    {
                        "code": code,
                        "family": family,
                        "type": r[col["TYPE"]] if "TYPE" in col else "",
                        "oem": r[col["OEM"]] if "OEM" in col else "ALL",
                        "hand": r[col["HAND"]] if "HAND" in col else "",
                        "width_in": int(width_m.group(0)) if width_m else None,
                        "description": r[col[desc_col]] if desc_col else "",
                        "price_cad": price,
                        "aliases": derive_aliases(code, family),
                    }
                )
        elif header[:3] == ["LINES", "INCENTIVE", "QUALIFYING STRUCTURE"]:
            for r in rows[1:]:
                if len(r) >= 3 and r[0]:
                    incentives.append({"lines": r[0], "incentive": r[1], "structure": r[2]})
        elif header[:2] == ["TERM", "COMMERCIAL BASIS"]:
            for r in rows[1:]:
                if len(r) >= 2 and r[0]:
                    terms.append({"term": r[0], "basis": r[1]})
        else:
            # Call-out boxes ("REQUIRED PAIRING …", "PO REQUIREMENT …", "SELECTION CHECKPOINT …") sit
            # in single-cell tables, or in the second row of the configuration-key table.
            for r in rows:
                for cell in r[:1]:
                    if cell.startswith(_NOTE_PREFIXES):
                        notes.append(cell)

    if not rows_out:
        raise PriceBookError(
            "No 'FTL ITEM CODE' / 'CATALOGUE PRICE' tables found — this does not look like the "
            "FTL Contractor Price Book."
        )
    families = {r["family"] for r in rows_out}
    missing = {"door_operator", "panel_adaptor", "clutch", "roller_guide", "governor"} - families
    if missing:
        raise PriceBookError(
            f"Price book parsed but is missing expected sections: {sorted(missing)}. "
            "Refusing to import a partial price book."
        )
    # de-duplicate notes, keep order
    seen_n = set()
    notes = [n for n in notes if not (n in seen_n or seen_n.add(n))]
    return {
        "revision": revision or "unknown revision",
        "rows": rows_out,
        "incentives": incentives,
        "terms": terms,
        "notes": notes,
    }


# ---------------------------------------------------------------------------------------------
# Overlay (consumed by rules_engine.py)
# ---------------------------------------------------------------------------------------------
SOURCE_LABEL = "FTL Contractor Price Book (Wittur Products 2026 Catalogue & Price List)"

# Conflict-register entries the price book settles. rules_engine.cite_conflict() renders these as
# RESOLVED instead of OPEN. Only I-001 is settled by a PRICE; every other open conflict (OL100 SKU,
# freight matrix, safety nomenclature, ...) is a different question this document does not answer.
RESOLVES_CONFLICTS = {
    "I-001": (
        "SGV2_DOOR_TOOLS is $698.50 (the old list's $598.50 is superseded; matches the "
        "client-confirmed EST-261132 figure). "
        "Whether the tool is bundled at $0 on a given project is still a per-project decision."
    )
}


def build_overlay(book: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "source": f"{SOURCE_LABEL}, {book['revision']}",
        "revision": book["revision"],
        "currency": "CAD",
        "rows": [
            {
                "code": r["code"],
                "family": r["family"],
                "type": r["type"],
                "oem": r["oem"],
                "hand": r["hand"],
                "width_in": r["width_in"],
                "description": r["description"],
                "price_cad": r["price_cad"],
                "aliases": r["aliases"],
            }
            for r in book["rows"]
        ],
        "incentives": book["incentives"],
        "resolves_conflicts": RESOLVES_CONFLICTS,
    }


def write_overlay(book: Dict[str, Any], path: str = OVERLAY_PATH) -> str:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        os.chmod(path, 0o644)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(build_overlay(book), f, indent=1, ensure_ascii=False)
        f.write("\n")
    return path


# ---------------------------------------------------------------------------------------------
# Index pages (consumed by pricelist_store.reindex_pricelist)
# ---------------------------------------------------------------------------------------------
def _money(v: float) -> str:
    return f"$ {v:,.2f}"


def _row_text(r: Dict[str, Any]) -> str:
    """One catalogue row in the same column order the old PDF index used (TYPE, O.E.M., PRODUCT
    NAME, DOOR HAND, DESCRIPTION, DOOR WIDTH, PRICE), skipping columns a section doesn't have."""
    if r["family"] == "door_tools":
        return f"ALL ALL {r['code']} ALL {r['description'].upper()} ALL {_money(r['price_cad'])}"
    parts = [r["type"] or "ALL", r["oem"] or "ALL", r["code"]]
    if r["hand"]:
        parts.append(r["hand"])
    parts.append(r["description"])
    if r["family"] in ("door_operator", "panel_adaptor") and r["width_in"] is not None:
        parts.append(str(r["width_in"]))
    parts.append(_money(r["price_cad"]))
    return " ".join(p for p in parts if p)


def _paginate(header: str, rows: List[str], trailer: str = "") -> List[str]:
    """Split rows across pages at ROW boundaries so no page exceeds PAGE_WORD_BUDGET words."""
    pages: List[str] = []
    cur: List[str] = []
    base = len(header.split()) + len(trailer.split())
    n = base
    for row in rows:
        w = len(row.split())
        if cur and n + w > PAGE_WORD_BUDGET:
            pages.append(" ".join([header] + cur + ([trailer] if trailer else [])))
            cur, n = [], base
        cur.append(row)
        n += w
    if cur:
        pages.append(" ".join([header] + cur + ([trailer] if trailer else [])))
    return pages


def load_legacy_blocks(path: str = LEGACY_BLOCKS_PATH) -> Dict[str, str]:
    if not os.path.exists(path):
        return {}
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return {b["id"]: b["text"] for b in data.get("blocks", [])}


def _split_legacy_panels(text: str) -> List[str]:
    """Universal-panel block -> per-row strings (header first), so it can be paginated at row
    boundaries like everything else."""
    parts = re.split(r"(?=\b(?:1S|2T|2C) (?:1S|2T|2C)_(?:T1_)?UNIVERSAL_CAR_DOOR_PANEL)", text)
    return [p.strip() for p in parts if p.strip()]


def build_pages(book: Dict[str, Any], legacy: Optional[Dict[str, str]] = None) -> List[str]:
    """Index pages for the whole catalogue: every row of the price book, then the retained legacy
    sections. Order is stable so unchanged pages hash-match and are not re-embedded."""
    legacy = load_legacy_blocks() if legacy is None else legacy
    rev = book["revision"]
    top = f"FTL DISTRIBUTION INC WITTUR PRODUCT CATALOGUE OFFERINGS (CONTRACTOR PRICE BOOK {rev} - CAD)"
    rows = book["rows"]
    by_family: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        by_family.setdefault(r["family"], []).append(r)

    pages: List[str] = []

    # 1. Door operators + programming tool (+ retained detector block)
    ops = by_family.get("door_operator", []) + by_family.get("door_tools", [])
    pages += _paginate(
        f"{top} SGV2 CAR DOOR OPERATOR - LINEAR SYSTEM MOD KITS (ALL OEMs) "
        "TYPE O.E.M. PRODUCT NAME DOOR HAND DESCRIPTION DOOR WIDTH CATALOGUE PRICE",
        [_row_text(r) for r in ops],
    )
    if legacy.get("detectors_and_notes"):
        pages.append(
            f"{top} VISIONPLUS DOOR DETECTION AND SYSTEM NOTES "
            "TYPE O.E.M. PRODUCT NAME DOOR HAND DESCRIPTION DOOR WIDTH CATALOGUE PRICE "
            + legacy["detectors_and_notes"]
        )

    # 2-4. Panel adaptors per door type
    adaptors = by_family.get("panel_adaptor", [])
    for dtype, label in (
        ("1S", "SINGLE SPEED, SIDE OPENING"),
        ("2T", "TWO SPEED, SIDE OPENING"),
        ("2C", "CENTER OPENING / BI-PARTING"),
    ):
        sel = [r for r in adaptors if r["type"] == dtype]
        if sel:
            pages += _paginate(
                f"{top} SGV2 DOOR COMPONENT KITS - OEM PANEL ADAPTORS ({dtype} - {label}) "
                "TYPE O.E.M. PRODUCT NAME DOOR HAND DESCRIPTION DOOR WIDTH CATALOGUE PRICE",
                [_row_text(r) for r in sel],
            )

    # 5. Clutches
    clutches = by_family.get("clutch", [])
    if clutches:
        pages += _paginate(
            f"{top} CLUTCH c/w CAR DOOR INTERLOCK (WITTUR COUPLER) - match OEM family and hand. "
            "Every SGV2 door operator purchase order must include a compatible clutch. "
            "TYPE O.E.M. PRODUCT NAME DOOR HAND DESCRIPTION CATALOGUE PRICE",
            [_row_text(r) for r in clutches],
        )

    # 6. Roller guides, governors, safeties (+ retained OL100 / swingarm rows)
    rg = by_family.get("roller_guide", [])
    gov = by_family.get("governor", []) + by_family.get("car_safety", [])
    legacy_motion = legacy.get("ol100_and_swingarm", "")
    if rg:
        pages += _paginate(
            f"{top} WITTUR ROLLER GUIDE ASSEMBLIES (CAR & CWT) - 4x PCS REQUIRED FOR FULL SET "
            "TYPE O.E.M. PRODUCT NAME DESCRIPTION CATALOGUE PRICE",
            [_row_text(r) for r in rg],
        )
    if gov or legacy_motion:
        pages += _paginate(
            f"{top} WITTUR MOTION & SAFETY GEAR - ELEVATOR GOVERNORS & UNIDIRECTIONAL SAFETIES "
            "(final selection subject to application review) "
            "TYPE O.E.M. PRODUCT NAME DESCRIPTION CATALOGUE PRICE",
            [_row_text(r) for r in gov] + ([legacy_motion] if legacy_motion else []),
        )

    # 7. Universal car door panels (retained)
    panels = legacy.get("universal_car_door_panels")
    if panels:
        parts = _split_legacy_panels(panels)
        pages += _paginate(f"{top} UNIVERSAL CAR DOOR PANELS", parts)

    # 8. Incentives + commercial terms (from the book)
    inc = book.get("incentives") or []
    if inc or book.get("terms"):
        bits = [f"{top} FTL BUNDLED PURCHASING INCENTIVE"]
        for i in inc:
            bits.append(f"{i['lines']} qualifying product lines: {i['incentive']} - {i['structure']}.")
        bits.append(
            "Bundled incentives apply to qualifying products only, require at least one primary "
            "system component, must be ordered on one purchase order, and are applied at invoicing. "
            "Not retroactive; may not be combined with other promotional or special pricing unless "
            "FTL approves otherwise in writing."
        )
        for n in book.get("notes") or []:
            bits.append(n)
        pages += _paginate(bits[0], bits[1:])
        term_rows = [f"{t['term']}: {t['basis']}" for t in book.get("terms") or []]
        if term_rows:
            pages += _paginate(f"{top} COMMERCIAL TERMS & CONDITIONS", term_rows)

    # 9. Spare parts (retained)
    spare = legacy.get("spare_parts")
    if spare:
        parts = re.split(r"(?=\b(?:WITTUR_SELCOM|SUPRA_MOTOR|MIDI_SUPRA|SVG2_BRACKET|FLAT_SCREW|SOCKET_CAP|CONTACT_WASHER))", spare)
        parts = [p.strip() for p in parts if p.strip()]
        pages += _paginate(f"{top} WITTUR SPARE PARTS (maintenance only)", parts)

    # Every page must fit in one chunk (see PAGE_WORD_BUDGET) — fail loudly, don't split a row.
    for i, p in enumerate(pages, 1):
        if len(p.split()) > 280:
            raise PriceBookError(f"Generated index page {i} is {len(p.split())} words (>280).")
    return pages


def price_book_summary(book: Dict[str, Any]) -> Dict[str, Any]:
    fam: Dict[str, int] = {}
    for r in book["rows"]:
        fam[r["family"]] = fam.get(r["family"], 0) + 1
    return {"revision": book["revision"], "rows": len(book["rows"]), "by_family": fam}
